# -*- coding: utf-8 -*-
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

D  = json.load(open('clinic_avg.json'))
R, T = D['rows'], D['total']
MONTHS = D['months']; QUARTERS = D['quarters']
AR = (sys.argv[1] if len(sys.argv)>1 else 'ar') == 'ar'
L  = D['label_ar'] if AR else D['label_en']

NAVY='1F3864'; BLUE='2E5FA3'; GREEN='1E7145'; GOLD='B07D00'; GREY='595959'; RED='9C0006'
F25='EAF1F8'; F26='E2EFDA'; XF='F2F2F2'
thin=Side(style='thin', color='BFBFBF'); BORD=Border(left=thin,right=thin,top=thin,bottom=thin)
wb=Workbook(); wb.remove(wb.active)
def t(a,e): return a if AR else e
ALIGN_TXT = 'right' if AR else 'left'

def sheet(name):
    ws=wb.create_sheet(name); ws.sheet_view.rightToLeft=AR; ws.sheet_view.showGridLines=False; return ws
def title_block(ws,title,sub,width):
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=width)
    c=ws.cell(1,1,title); c.font=Font(bold=True,size=16,color='FFFFFF')
    c.fill=PatternFill('solid',fgColor=NAVY); c.alignment=Alignment(horizontal='center',vertical='center')
    ws.row_dimensions[1].height=30
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=width)
    c=ws.cell(2,1,sub); c.font=Font(size=10,color=GREY)
    c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True); ws.row_dimensions[2].height=34
def header_row(ws,r,headers,widths,fill=BLUE):
    for i,(h,w) in enumerate(zip(headers,widths),start=1):
        c=ws.cell(r,i,h); c.font=Font(bold=True,size=10,color='FFFFFF')
        c.fill=PatternFill('solid',fgColor=fill); c.border=BORD
        c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width=w
    ws.row_dimensions[r].height=36; ws.freeze_panes=ws.cell(r+1,3)

SUB = t('متوسط الشراء لكل عيادة · %d عيادة · %d طلب · %s د.ك · من 01/2025 إلى 24/09/2026 (21 شهراً)',
        'Average purchasing per clinic · %d clinics · %d orders · %s KD · 01/2025 to 24/09/2026 (21 months)')

# ============================================================ 1) SUMMARY
ws=sheet(t('الملخص','Summary'))
title_block(ws, t('UltraMed — متوسط شراء العيادات 2025–2026','UltraMed — Clinic Purchasing Averages 2025–2026'),
    SUB % (T['accounts'], T['orders'], format(round(T['net']),',')), 8)
ws.column_dimensions['A'].width=3
for col,w in zip('BCDEFGH',[30,16,16,16,16,16,30]): ws.column_dimensions[col].width=w
r=4
def h2(x):
    global r
    ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=8)
    c=ws.cell(r,2,x); c.font=Font(bold=True,size=12,color='FFFFFF')
    c.fill=PatternFill('solid',fgColor=BLUE); c.alignment=Alignment(horizontal=ALIGN_TXT,vertical='center')
    ws.row_dimensions[r].height=22; r+=1
def para(x,bold=False):
    global r
    ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=8)
    c=ws.cell(r,2,x); c.font=Font(size=10,bold=bold)
    c.alignment=Alignment(horizontal=ALIGN_TXT,vertical='top',wrap_text=True)
    ws.row_dimensions[r].height=max(16,15*(len(x)//115+1)); r+=1
def tbl(headers,rows,fmts=None,fills=None,first_col=2):
    global r
    for j,h in enumerate(headers,start=first_col):
        c=ws.cell(r,j,h); c.font=Font(bold=True,size=10,color='FFFFFF')
        c.fill=PatternFill('solid',fgColor=NAVY); c.border=BORD
        c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    ws.row_dimensions[r].height=26; r+=1
    for i,row in enumerate(rows):
        for j,v in enumerate(row,start=first_col):
            c=ws.cell(r,j,v); c.border=BORD; c.font=Font(size=10,bold=(j==first_col))
            c.alignment=Alignment(horizontal=ALIGN_TXT if j==first_col else 'center',vertical='center',wrap_text=True)
            if fmts and fmts.get(j-first_col+1): c.number_format=fmts[j-first_col+1]
            if fills: c.fill=PatternFill('solid',fgColor=fills[i])
        ws.row_dimensions[r].height=20; r+=1
    r+=1

h2(t('الخلاصة','The headline'))
para(t('%d عيادة اشترت فعلياً منذ انطلاق القناة، بإجمالي %s د.ك على %d طلب. متوسط الطلب %s د.ك.'
       % (T['accounts'], format(round(T['net']),','), T['orders'], format(round(T['net']/T['orders']),',')),
       '%d clinics have purchased since the channel launched, totalling %s KD across %d orders. Average order %s KD.'
       % (T['accounts'], format(round(T['net']),','), T['orders'], format(round(T['net']/T['orders']),','))), bold=True)
para(t('قناة العيادات بدأت في مايو 2025 — لا يوجد أي حساب طبي في 2024 تحت أي فئة، ولا في يناير–أبريل 2025. هذه الحقيقة تحكم كل مقارنة سنوية في هذا الملف.',
       'The clinic channel began in May 2025 — there is no medical account in 2024 under any class, nor in January–April 2025. That fact governs every year-on-year comparison in this file.'))

h2(t('النمو على نوافذ زمنية متطابقة','Growth on identical windows'))
para(t('لأن البيانات تنتهي في 24/09/2026، تُقارن كل نافذة بنفس الأيام بالضبط من السنة السابقة — بلا أي تقدير أو تعديل حسابي.',
       'Because the data ends on 24/09/2026, each window is compared against exactly the same days of the previous year — with no estimate or adjustment.'))
tbl([t('النافذة','Window'), '2025 (KD)', '2026 (KD)', t('النمو','Growth'), t('الأيام','Days'), t('قراءتها','How to read it')],
    [[t('01/05 → 24/09  (الأساسية)','01/05 → 24/09  (primary)'), T['win25'], T['win26'], T['growth'], 147,
      t('المقارنة الصحيحة: أول شهر فيه نشاط في 2025 هو مايو','The correct comparison: May is the first active month of 2025')],
     [t('01/07 → 24/09  (الربع الثالث)','01/07 → 24/09  (Q3)'), T['q325'], T['q326'], T['q3_growth'], 86,
      t('الربع الثالث مقابل نظيره، 86 يوماً لكل سنة','Q3 against its counterpart, 86 days each year')],
     [t('01/01 → 24/09  (مضلِّلة)','01/01 → 24/09  (misleading)'), T['ytd25'], T['ytd26'], T['ytd_growth'], 267,
      t('لا تُستخدم: يناير–أبريل 2025 أصفار فيتضخم النمو','Do not use: Jan–Apr 2025 are zeros, so growth is inflated')]],
    {2:'#,##0',3:'#,##0',4:'+0.0%;-0.0%',5:'#,##0'}, [F26,F26,XF])
para(t('أي أن النمو الحقيقي للقناة +%.0f%% على النافذة الأساسية، و+%.0f%% في الربع الثالث — لا +%.0f%% كما تُظهر المقارنة من أول السنة.'
       % (T['growth']*100, T['q3_growth']*100, T['ytd_growth']*100),
       'So the channel\'s real growth is +%.0f%% on the primary window and +%.0f%% in Q3 — not the +%.0f%% that a year-to-date comparison shows.'
       % (T['growth']*100, T['q3_growth']*100, T['ytd_growth']*100)), bold=True)

h2(t('الشراء الشهري — كل العيادات','Monthly purchasing — all clinics'))
tbl([t('الشهر','Month'), t('الشراء KD','Purchases KD'), t('الطلبات','Orders'),
     t('عيادات نشطة','Active clinics'), t('متوسط الطلب KD','Avg order KD'), t('متوسط العيادة KD','Avg per clinic KD')],
    [[L[i], T['monthly'][m], T['orders_monthly'][m], T['active_accounts'][m],
      round(T['monthly'][m]/T['orders_monthly'][m],1) if T['orders_monthly'][m] else 0,
      round(T['monthly'][m]/T['active_accounts'][m],1) if T['active_accounts'][m] else 0] for i,m in enumerate(MONTHS)],
    {2:'#,##0',3:'#,##0',4:'#,##0',5:'#,##0',6:'#,##0'},
    [F25 if m.startswith('2025') else F26 for m in MONTHS])

h2(t('الشراء الربعي','Quarterly purchasing'))
tbl([t('الربع','Quarter'), t('الشراء KD','Purchases KD'), t('الأيام المغطاة','Days covered'), t('ملاحظة','Note')],
    [[k, T['q'][k], dy,
      (t('ناقص: %d من 92 يوماً — يُقارن بنظيره لا بالربع السابق','Partial: %d of 92 days — compare with its counterpart, not the previous quarter') % dy)
      if k in D['q_full'] else (t('القناة لم تبدأ بعد','Channel had not started') if T['q'][k]==0 else t('مكتمل','Complete'))]
     for k,dy in QUARTERS],
    {2:'#,##0',3:'#,##0'}, [F25 if k.startswith('2025') else F26 for k,_ in QUARTERS])

h2(t('دورة حياة العيادات','Clinic lifecycle'))
from collections import Counter
life = Counter(x['life_ar'] if AR else x['life_en'] for x in R)
order_ar=['نامية','مستقرة','عادت في نافذة 2026','جديدة في 2026','متراجعة','مفقودة — لم تشترِ في 2026','إرجاع كامل — لم يتحول إلى شراء']
order_en=['Growing','Stable','Returned in the 2026 window','New in 2026','Declining','Lost — no purchase in 2026','Fully returned — never converted']
keys = order_ar if AR else order_en
tbl([t('الحالة','Status'), t('عدد العيادات','Clinics'), t('قيمتها 2026 KD','Their 2026 value KD'), t('ماذا تعني','What it means')],
    [[k, life.get(k,0),
      round(sum(x['win26'] for x in R if (x['life_ar'] if AR else x['life_en'])==k),1),
      {'نامية':'شراؤها في 2026 أعلى 10%+ من نفس نافذة 2025','مستقرة':'شراؤها ضمن ±10% من السنة السابقة',
       'عادت في نافذة 2026':'لم تشترِ في نافذة 2025 (مايو–سبتمبر) وعادت في نافذة 2026','جديدة في 2026':'أول طلب لها في 2026',
       'متراجعة':'شراؤها في 2026 أقل 10%+ من نفس نافذة 2025','مفقودة — لم تشترِ في 2026':'اشترت في 2025 ولم تشترِ في 2026 — أولوية استرجاع',
       'Growing':'2026 purchases 10%+ above the same 2025 window','Stable':'Within ±10% of last year',
       'Returned in the 2026 window':'No purchase in the 2025 window, back in the 2026 one','New in 2026':'First order falls in 2026',
       'Declining':'2026 purchases 10%+ below the same 2025 window','Lost — no purchase in 2026':'Bought in 2025, nothing in 2026 — win-back priority',
       'إرجاع كامل — لم يتحول إلى شراء':'استلمت بضاعة وأرجعتها كاملة — عرض/أمانة لا شراء، تُقرأ منفصلة عن العيادات',
       'Fully returned — never converted':'Received stock and returned all of it — consignment, not a purchase; read separately'}.get(k,'')]
     for k in keys],
    {2:'#,##0',3:'#,##0'},
    [F26,F26,F26,F26,'FFF2CC','FCE4D6','EDEDED'])

lost=[x for x in R if (x['life_ar'] if AR else x['life_en']).startswith(t('مفقودة','Lost'))]
if lost:
    h2(t('الأولوية العملية — عيادات مفقودة','The practical priority — lost clinics'))
    para(t('%d عيادة اشترت في 2025 ولم تشترِ شيئاً في 2026 كلها، بقيمة %s د.ك ضائعة. استرجاع عميل سابق أسرع من فتح حساب جديد.'
           % (len(lost), format(round(sum(x['net25'] for x in lost)),',')),
           '%d clinics bought in 2025 and nothing in all of 2026, worth %s KD lost. Winning back a former customer is faster than opening a new account.'
           % (len(lost), format(round(sum(x['net25'] for x in lost)),','))), bold=True)
    tbl([t('العيادة','Clinic'), t('شراء 2025 كاملة KD','Full-year 2025 purchases KD'), t('طلبات','Orders'),
         t('آخر طلب','Last order'), t('أيام بلا طلب','Days silent'), t('متوسط الطلب KD','Avg order KD')],
        [[x['account'], x['net25'], x['orders'], x['last'], x['silent'], x['avg_order']]
         for x in sorted(lost, key=lambda x:-x['win25'])],
        {2:'#,##0',3:'#,##0',5:'#,##0',6:'#,##0'}, ['FCE4D6']*len(lost))

# ============================================================ 2) MONTHLY AVERAGE
ws=sheet(t('المتوسط الشهري','Monthly Average'))
H=[t('#','#'),t('العيادة','Clinic'),t('النوع','Type'),t('المندوب','Rep'),t('أول طلب','First order'),t('آخر طلب','Last order'),
   t('أيام بلا طلب','Days silent'),t('الطلبات','Orders'),t('إجمالي الشراء KD','Total purchases KD'),
   t('أشهر نشطة','Active months'),t('أشهر منذ أول طلب','Months since first order'),
   t('متوسط/شهر نشط KD','Avg per active month KD'),t('متوسط/شهر معياري KD','Avg per standardised month KD'),
   t('متوسط الطلب KD','Avg order KD'),t('طلبات/شهر','Orders/month'),
   t('براندات','Brands'),t('أصناف','SKUs'),t('قطع مدفوعة','Paid units'),t('قطع مجانية','Free units'),
   t('مرتجعات KD','Returns KD'),t('الحالة','Status'),t('دورة الحياة','Lifecycle')]
W=[5,44,12,18,12,12,11,9,15,10,13,15,16,13,11,9,9,11,11,12,15,24]
title_block(ws, t('المتوسط الشهري لكل عيادة','Monthly average by clinic'),
    t('«متوسط/شهر نشط» = الإجمالي ÷ الأشهر التي فيها طلب — ماذا تشتري العيادة حين تشتري  ·  «متوسط/شهر معياري» = الإجمالي ÷ الأشهر من أول طلب حتى 24/09/2026 — يعاقب الانقطاع. الفرق بين العمودين هو مقياس انتظام العيادة.',
      '"Avg per active month" = total ÷ months with an order — what the clinic buys when it buys  ·  "Avg per standardised month" = total ÷ months since its first order to 24/09/2026 — this penalises silence. The gap between the two columns measures how regular the clinic is.'), len(H))
header_row(ws,3,H,W)
r=4
for i,x in enumerate(R,start=1):
    vals=[i,x['account'],x['cls_ar'] if AR else x['cls_en'],x['reps_ar'] if AR else x['reps_en'],
          x['first'],x['last'],x['silent'],x['orders'],x['net'],x['active_months'],x['span_months'],
          x['avg_active'],x['avg_span'],x['avg_order'],x['orders_per_month'],x['brands'],x['skus'],
          x['units'],x['foc'],x['ret'],x['status_ar'] if AR else x['status_en'],x['life_ar'] if AR else x['life_en']]
    for j,v in enumerate(vals,start=1):
        c=ws.cell(r,j,v); c.border=BORD; c.font=Font(size=9)
        c.alignment=Alignment(horizontal=ALIGN_TXT if j in (2,22) else 'center',vertical='center',wrap_text=(j in (2,22)))
        if j in (9,12,13,14,20): c.number_format='#,##0.0'
        if j in (7,8,10,16,17,18,19): c.number_format='#,##0'
        if j in (11,15): c.number_format='#,##0.00'
        if j==21: c.font=Font(size=9,bold=True,color={'نشطة':GREEN,'Active':GREEN,'تحتاج زيارة':GOLD,'Needs a visit':GOLD,
                                                      'منقطعة':RED,'Lapsed':RED,'طلب واحد فقط':GREY,'One order only':GREY}[v])
        if j==22: c.font=Font(size=9,bold=True,color=(RED if 'مفقودة' in str(v) or 'Lost' in str(v)
                                                      else GOLD if 'متراجعة' in str(v) or 'Declining' in str(v)
                                                      else GREEN if 'نامية' in str(v) or 'Growing' in str(v) else GREY))
    ws.row_dimensions[r].height=22; r+=1
for j in range(1,len(H)+1):
    c=ws.cell(r,j); c.border=BORD; c.fill=PatternFill('solid',fgColor=NAVY)
    c.font=Font(bold=True,size=10,color='FFFFFF'); c.alignment=Alignment(horizontal='center',vertical='center')
ws.cell(r,2,t('الإجمالي','Total'))
ws.cell(r,8,T['orders']).number_format='#,##0'; ws.cell(r,9,T['net']).number_format='#,##0'
ws.cell(r,14,round(T['net']/T['orders'],1)).number_format='#,##0.0'
ws.auto_filter.ref='A3:%s%d'%(get_column_letter(len(H)),r-1)

# ============================================================ 3) MONTHLY DETAIL
ws=sheet(t('التفصيل الشهري','Monthly Detail'))
H=[t('#','#'),t('العيادة','Clinic')]+L+[t('الإجمالي','Total')]
W=[5,42]+[11]*len(MONTHS)+[13]
title_block(ws, t('التفصيل الشهري — الشراء بالدينار لكل عيادة في كل شهر','Monthly detail — purchases in KD per clinic per month'),
    t('الخلية الفارغة (—) تعني شهراً بلا أي طلب. أشهر 2025 بخلفية زرقاء و2026 بخلفية خضراء · سبتمبر 2026 حتى يوم 24 فقط.',
      'An empty cell (—) means a month with no order. 2025 months are shaded blue and 2026 green · September 2026 covers only up to the 24th.'), len(H))
header_row(ws,3,H,W)
for j,m in enumerate(MONTHS, start=3):
    ws.cell(3,j).fill=PatternFill('solid',fgColor=('4A7EBB' if m.startswith('2025') else '4E8A52'))
r=4
for i,x in enumerate(R,start=1):
    ws.cell(r,1,i).border=BORD; ws.cell(r,1).font=Font(size=9); ws.cell(r,1).alignment=Alignment(horizontal='center')
    c=ws.cell(r,2,x['account']); c.border=BORD; c.font=Font(size=9); c.alignment=Alignment(horizontal=ALIGN_TXT,vertical='center')
    for j,m in enumerate(MONTHS,start=3):
        v=x['monthly'][m]
        c=ws.cell(r,j, v if v else '—'); c.border=BORD; c.font=Font(size=9)
        c.alignment=Alignment(horizontal='center',vertical='center')
        if v: c.number_format='#,##0'
        c.fill=PatternFill('solid',fgColor=(F25 if m.startswith('2025') else F26)) if v else PatternFill('solid',fgColor='FAFAFA')
    c=ws.cell(r,len(H),x['net']); c.border=BORD; c.font=Font(size=9,bold=True); c.number_format='#,##0'
    c.alignment=Alignment(horizontal='center',vertical='center')
    ws.row_dimensions[r].height=20; r+=1
for j in range(1,len(H)+1):
    c=ws.cell(r,j); c.border=BORD; c.fill=PatternFill('solid',fgColor=NAVY); c.font=Font(bold=True,size=9,color='FFFFFF')
    c.alignment=Alignment(horizontal='center',vertical='center')
ws.cell(r,2,t('الإجمالي','Total'))
for j,m in enumerate(MONTHS,start=3): ws.cell(r,j,T['monthly'][m]).number_format='#,##0'
ws.cell(r,len(H),T['net']).number_format='#,##0'
ws.auto_filter.ref='A3:%s%d'%(get_column_letter(len(H)),r-1)

# ============================================================ 4) QUARTERS
ws=sheet(t('الأرباع','Quarters'))
QK=[k for k,_ in QUARTERS]
H=[t('#','#'),t('العيادة','Clinic')]+QK+[t('الإجمالي','Total'),t('الربع 3 / 2025 (86 يوم)','Q3 2025 (86 days)'),
   t('الربع 3 / 2026 (86 يوم)','Q3 2026 (86 days)'),t('نمو الربع 3','Q3 growth')]
W=[5,42]+[12]*len(QK)+[13,16,16,13]
title_block(ws, t('الشراء الربعي لكل عيادة','Quarterly purchasing by clinic'),
    t('الربع الثالث 2026 مغطى 86 يوماً من 92، فالعمودان الأخيران يقارنانه بنفس الـ86 يوماً من 2025 — مقارنة متطابقة بلا تقدير. الربع الأول 2025 صفر لأن القناة بدأت في مايو.',
      'Q3 2026 covers 86 of 92 days, so the last two columns compare it against the same 86 days of 2025 — an identical comparison with no estimate. Q1 2025 is zero because the channel started in May.'), len(H))
header_row(ws,3,H,W)
for j,k in enumerate(QK,start=3):
    ws.cell(3,j).fill=PatternFill('solid',fgColor=('4A7EBB' if k.startswith('2025') else '4E8A52'))
r=4
for i,x in enumerate(R,start=1):
    vals=[i,x['account']]+[x['q'][k] or '—' for k in QK]+[x['net'], x['q325'] or '—', x['q326'] or '—',
          x['q3_growth'] if x['q3_growth'] is not None else '—']
    for j,v in enumerate(vals,start=1):
        c=ws.cell(r,j,v); c.border=BORD; c.font=Font(size=9)
        c.alignment=Alignment(horizontal=ALIGN_TXT if j==2 else 'center',vertical='center')
        if isinstance(v,(int,float)) and j>2: c.number_format='#,##0'
        if 3<=j<=2+len(QK): c.fill=PatternFill('solid',fgColor=(F25 if QK[j-3].startswith('2025') else F26)) if v!='—' else PatternFill('solid',fgColor='FAFAFA')
        if j==len(H):
            if isinstance(v,float):
                c.number_format='+0%;-0%'; c.font=Font(size=9,bold=True,color=(GREEN if v>0 else RED))
        if j==3+len(QK): c.font=Font(size=9,bold=True)
    ws.row_dimensions[r].height=20; r+=1
for j in range(1,len(H)+1):
    c=ws.cell(r,j); c.border=BORD; c.fill=PatternFill('solid',fgColor=NAVY); c.font=Font(bold=True,size=9,color='FFFFFF')
    c.alignment=Alignment(horizontal='center',vertical='center')
ws.cell(r,2,t('الإجمالي','Total'))
for j,k in enumerate(QK,start=3): ws.cell(r,j,T['q'][k]).number_format='#,##0'
ws.cell(r,3+len(QK),T['net']).number_format='#,##0'
ws.cell(r,4+len(QK),T['q325']).number_format='#,##0'; ws.cell(r,5+len(QK),T['q326']).number_format='#,##0'
c=ws.cell(r,6+len(QK),T['q3_growth']); c.number_format='+0%;-0%'
ws.auto_filter.ref='A3:%s%d'%(get_column_letter(len(H)),r-1)

# ============================================================ 5) YEAR ON YEAR
ws=sheet(t('المقارنة السنوية','Year on Year'))
H=[t('#','#'),t('العيادة','Clinic'),t('النوع','Type'),
   t('01/05→24/09 من 2025 KD','01/05→24/09 of 2025 KD'),t('01/05→24/09 من 2026 KD','01/05→24/09 of 2026 KD'),
   t('الفرق KD','Change KD'),t('النمو','Growth'),
   t('طلبات 2025','2025 orders'),t('طلبات 2026','2026 orders'),
   t('2025 كاملة KD','Full 2025 KD'),t('2026 حتى 24/09 KD','2026 to 24/09 KD'),
   t('متوسط الطلب KD','Avg order KD'),t('دورة الحياة','Lifecycle'),t('الإجراء المقترح','Suggested action')]
W=[5,42,12,18,18,13,11,11,11,14,16,13,24,46]
title_block(ws, t('المقارنة السنوية على نافذة متطابقة','Year-on-year on an identical window'),
    t('النافذة 01/05 → 24/09 من كل سنة (147 يوماً) — اختيرت لأن قناة العيادات بدأت في مايو 2025، فالمقارنة من أول السنة تُضخّم النمو إلى +%.0f%% بدل +%.0f%% الحقيقي.'
      % (T['ytd_growth']*100, T['growth']*100),
      'The window is 01/05 → 24/09 of each year (147 days) — chosen because the clinic channel started in May 2025, so a year-to-date comparison inflates growth to +%.0f%% instead of the real +%.0f%%.'
      % (T['ytd_growth']*100, T['growth']*100)), len(H))
header_row(ws,3,H,W)
ACT = {
 'مفقودة — لم تشترِ في 2026':'أولوية أولى: زيارة استرجاع — كانت عميلاً فعلياً',
 'متراجعة':'زيارة تشخيص: ما سبب التراجع — سعر، منافس، أم انقطاع تواصل؟',
 'نامية':'تثبيت: عرض المرحلة 2 وتوسيع الرف',
 'مستقرة':'محاولة تنمية: صنف جديد من المرحلة 2',
 'عادت في نافذة 2026':'متابعة قريبة حتى يثبت الطلب الثالث',
 'جديدة في 2026':'متابعة حتى إعادة الطلب — الطلب الثاني هو المؤشر',
 'Lost — no purchase in 2026':'First priority: win-back visit — this was a real customer',
 'Declining':'Diagnostic visit: why the decline — price, competitor, or lost contact?',
 'Growing':'Consolidate: present Phase 2 and widen the shelf',
 'Stable':'Attempt growth: one new item from Phase 2',
 'Returned in the 2026 window':'Close follow-up until a third order lands',
 'New in 2026':'Follow up to a reorder — the second order is the signal'}
r=4
for i,x in enumerate(sorted(R,key=lambda x:(0 if (x['net26']==0 and x['net25']>0) else 1, -abs(x['win26']-x['win25']))),start=1):
    lf=x['life_ar'] if AR else x['life_en']
    vals=[i,x['account'],x['cls_ar'] if AR else x['cls_en'],x['win25'] or '—',x['win26'] or '—',
          round(x['win26']-x['win25'],1), x['growth'] if x['growth'] is not None else '—',
          x['win25_orders'] or '—',x['win26_orders'] or '—',x['net25'] or '—',x['net26'] or '—',x['avg_order'],lf,ACT.get(lf,'—')]
    for j,v in enumerate(vals,start=1):
        c=ws.cell(r,j,v); c.border=BORD; c.font=Font(size=9)
        c.alignment=Alignment(horizontal=ALIGN_TXT if j in (2,13,14) else 'center',vertical='center',wrap_text=(j in (2,13,14)))
        if j in (4,5,6,10,11,12) and isinstance(v,(int,float)): c.number_format='#,##0.0'
        if j==7 and isinstance(v,float):
            c.number_format='+0%;-0%'; c.font=Font(size=9,bold=True,color=(GREEN if v>0 else RED))
        if j==6 and isinstance(v,(int,float)): c.font=Font(size=9,bold=True,color=(GREEN if v>0 else RED if v<0 else GREY))
        if j==13: c.font=Font(size=9,bold=True,color=(RED if 'مفقودة' in str(v) or 'Lost' in str(v)
                                                      else GOLD if 'متراجعة' in str(v) or 'Declining' in str(v)
                                                      else GREEN if 'نامية' in str(v) or 'Growing' in str(v) else GREY))
    ws.row_dimensions[r].height=24; r+=1
ws.auto_filter.ref='A3:%s%d'%(get_column_letter(len(H)),r-1)

# ============================================================ 6) METHOD
ws=sheet(t('الطريقة','Method'))
title_block(ws, t('الطريقة والمعايير','Method and criteria'),
    t('كل رقم في هذا الملف محسوب بقاعدة مكتوبة قابلة لإعادة التطبيق على أي تصدير لاحق','Every figure in this file is computed by a written rule that can be re-applied to any later export'), 8)
ws.column_dimensions['A'].width=3
for col,w in zip('BCDEFGH',[34,22,22,22,22,22,28]): ws.column_dimensions[col].width=w
r=4
h2(t('1 — المصدر والنطاق','1 — Source and scope'))
para(t('ملف ERP الخام (MonthlySalesSummary) من 02/01/2024 إلى 24/09/2026 — 39,027 سطر فاتورة ومرتجع لكل المندوبين والقنوات.',
       'The raw ERP file (MonthlySalesSummary) from 02/01/2024 to 24/09/2026 — 39,027 invoice and return lines across all reps and channels.'))
para(t('النطاق: الحسابات المصنّفة Clinics و Hospital و Doctors و Government — 2,004 سطراً، كلها من فواتير رانوفا أيمن ومريم زهير حصراً (لا يخدم العيادات مندوب آخر في أي سنة).',
       'Scope: accounts classified as Clinics, Hospital, Doctors and Government — 2,004 lines, all of them invoiced by Ranova Ayman and Mariam Zohair exclusively (no other rep serves clinics in any year).'))
para(t('تحقّق مهم: لا يوجد أي حساب طبي في 2024 تحت أي فئة، ولا حساب غيّر فئته بين السنوات، ولا حساب يبدو اسمه عيادة وهو مصنّف «Other Customers». لذلك 2024 مستبعدة كلياً من هذا الملف — القناة بدأت في مايو 2025.',
       'An important check: there is no medical account in 2024 under any class, no account changed class between years, and no account with a clinic-like name sits under "Other Customers". 2024 is therefore excluded entirely from this file — the channel began in May 2025.'), bold=True)
para(t('تحقّق ثانٍ: أرقام العيادات في هذا التصدير تطابق التصدير السابق حرفياً للفترة المشتركة (998 سطراً · 81,263.2 د.ك · 47 حساباً)، فلا انحراف بين الملفين.',
       'A second check: the clinic figures in this export match the previous export exactly for the overlapping period (998 lines · 81,263.2 KD · 47 accounts), so there is no drift between the two files.'))
gap = lambda: None
r+=1
h2(t('2 — تعريف الشراء','2 — What counts as a purchase'))
tbl([t('المصطلح','Term'),t('التعريف','Definition')],
    [[t('الشراء','Purchase'), t('صافي المبيعات بالدينار: الفواتير ناقص المرتجعات. لا يشمل البضاعة المجانية لأن قيمتها صفر.','Net sales in KD: invoices minus returns. Free goods are excluded because their value is zero.')],
     [t('الطلب','Order'), t('فاتورة واحدة لها قيمة (Net Sales ≠ 0). فواتير القيمة الصفرية بونص وليست طلباً.','A single invoice carrying value (Net Sales ≠ 0). Zero-value invoices are bonus stock, not an order.')],
     [t('قطع مجانية','Free units'), t('كمية سطور الفواتير الصفرية — تُعرض في عمود مستقل لأنها الخصم الفعلي الممنوح للعيادة.','Quantity on zero-value invoice lines — shown in its own column because it is the real discount given to the clinic.')],
     [t('شهر نشط','Active month'), t('شهر فيه طلب واحد على الأقل بقيمة.','A month with at least one order carrying value.')]],
    {}, None)
h2(t('3 — المتوسطان الشهريان ولماذا هما اثنان','3 — The two monthly averages, and why there are two'))
tbl([t('المتوسط','Average'),t('الحساب','Calculation'),t('ما يقيسه','What it measures')],
    [[t('متوسط/شهر نشط','Per active month'), t('الإجمالي ÷ عدد الأشهر التي فيها طلب','Total ÷ months with an order'),
      t('حجم سلة العيادة حين تشتري — يُستخدم لتقدير الأوردر المتوقع','The size of the clinic\'s basket when it buys — used to size an expected order')],
     [t('متوسط/شهر معياري','Per standardised month'), t('الإجمالي ÷ (الأيام من أول طلب حتى 24/09/2026 ÷ 30.44)','Total ÷ (days from first order to 24/09/2026 ÷ 30.44)'),
      t('القيمة الشهرية الحقيقية للعيادة — يعاقب الانقطاع ويُستخدم للترتيب','The clinic\'s true monthly worth — it penalises silence and is used for ranking')]],
    {}, None)
para(t('الفرق بين العمودين هو مقياس الانتظام: عيادة متوسطها النشط 6,600 د.ك ومعياريها 1,786 د.ك اشترت مرتين فقط بمبلغ كبير — ليست عميلاً شهرياً.',
       'The gap between the two columns measures regularity: a clinic averaging 6,600 KD per active month but 1,786 KD per standardised month bought twice in large amounts — it is not a monthly customer.'), bold=True)
h2(t('4 — الأرباع والمقارنة السنوية','4 — Quarters and the year-on-year comparison'))
para(t('الأرباع تقويمية. الربع الثالث 2026 مغطى 86 يوماً من 92 لأن البيانات تنتهي في 24/09 — ولم أُعدّله تقديرياً، بل قارنته بنفس الـ86 يوماً من 2025، وهي مقارنة متطابقة لا تحتمل الخطأ.',
       'Quarters are calendar quarters. Q3 2026 covers 86 of 92 days because the data ends on 24/09 — and rather than adjusting it by estimate, it is compared against the same 86 days of 2025, an identical comparison that cannot mislead.'))
para(t('النافذة السنوية الأساسية 01/05 → 24/09 لأن القناة بدأت في مايو 2025. المقارنة من أول السنة تُظهر نمواً +%.0f%% وهو رقم مضلِّل، والحقيقي +%.0f%%. العمودان معروضان في ورقة الملخص مع وسم كل منهما.'
       % (T['ytd_growth']*100, T['growth']*100),
       'The primary annual window is 01/05 → 24/09 because the channel began in May 2025. A year-to-date comparison shows +%.0f%% growth, which is misleading; the real figure is +%.0f%%. Both are shown on the summary sheet, each labelled.'
       % (T['ytd_growth']*100, T['growth']*100)), bold=True)
h2(t('5 — تصنيف الحالة ودورة الحياة','5 — Status and lifecycle classification'))
tbl([t('التصنيف','Class'),t('القاعدة','Rule')],
    [[t('نشطة','Active'), t('آخر طلب خلال 45 يوماً','Last order within 45 days')],
     [t('تحتاج زيارة','Needs a visit'), t('آخر طلب بين 46 و90 يوماً','Last order between 46 and 90 days')],
     [t('منقطعة','Lapsed'), t('آخر طلب قبل أكثر من 90 يوماً','Last order more than 90 days ago')],
     [t('طلب واحد فقط','One order only'), t('لم تُعِد الطلب مطلقاً — يتقدم على أي تصنيف آخر','Never reordered — this overrides any other class')],
     [t('نامية / متراجعة','Growing / Declining'), t('فرق يزيد على ±10% بين نافذتي 2025 و2026','A difference of more than ±10% between the 2025 and 2026 windows')],
     [t('مفقودة','Lost'), t('شراء في 2025 وصفر في 2026 كلها — لا في نافذة المقارنة فقط','Purchases in 2025 and zero across all of 2026 — not merely in the comparison window')],
     [t('إرجاع كامل','Fully returned'), t('صافي الشراء صفر أو سالب لأن المرتجعات ألغت الفواتير — بضاعة عرض/أمانة لا شراء','Net purchases zero or negative because returns cancelled the invoices — consignment stock, not a purchase')]],
    {}, None)
h2(t('6 — حدود هذه القراءة','6 — Limits of this reading'))
para(t('•  سبتمبر 2026 حتى يوم 24 فقط — أي متوسط شهري يشمله محسوب على شهر ناقص 6 أيام.',
       '•  September 2026 runs only to the 24th — any monthly average including it rests on a month short by 6 days.'))
para(t('•  القناة عمرها 17 شهر نشاط فقط (05/2025 → 09/2026)، فالموسمية السنوية لا تزال غير مؤكدة: لدينا رمضان واحد وصيف واحد كاملان للمقارنة.',
       '•  The channel is only 17 active months old (05/2025 → 09/2026), so annual seasonality remains unconfirmed: we have one complete Ramadan and one complete summer to compare.'))
para(t('•  الأرقام تخص حسابات مصنّفة طبياً فقط. أي عيادة تشتري عبر الأونلاين أو عبر حساب موزّع لا تظهر هنا.',
       '•  The figures cover medically classified accounts only. A clinic buying through the online channel or a distributor account does not appear here.'))
para(t('•  وزارة الصحة حساب استثنائي (طلبان بمتوسط 6,600 د.ك) — يُقرأ منفصلاً عن العيادات الخاصة في أي متوسط عام.',
       '•  The Ministry of Health is an exceptional account (two orders averaging 6,600 KD) — read it separately from private clinics in any overall average.'))

fn = 'UltraMed-Clinic-Purchase-Averages-2025-2026%s.xlsx' % ('' if AR else '-EN')
wb.save(fn); print('saved', fn)
